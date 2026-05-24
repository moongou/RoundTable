from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT_PATH = REPO_ROOT / "backend" / "scripts" / "generate_ai_kids_lesson_pptx.py"
SOURCE_PATH = REPO_ROOT / "docs" / "小学四年级人工智能奇妙之旅_讲稿.md"


def test_generate_ai_kids_lesson_pptx_creates_editable_text(tmp_path: Path) -> None:
    output_path = tmp_path / "ai-kids-lesson.pptx"

    subprocess.run(
        [sys.executable, str(SCRIPT_PATH), "--source", str(SOURCE_PATH), "--output", str(output_path)],
        cwd=REPO_ROOT / "backend",
        check=True,
    )

    presentation = Presentation(output_path)
    assert len(presentation.slides) == 60

    first_slide = presentation.slides[0]
    first_slide_text = "\n".join(
        shape.text.strip()
        for shape in first_slide.shapes
        if hasattr(shape, "text") and shape.text and shape.text.strip()
    )

    assert "你好，人工智能！" in first_slide_text
    assert "给小学四年级小朋友的一堂奇妙科学课" in first_slide_text
    assert "让孩子理解：你好，人工智能！。" in first_slide_text
    assert "你最近在哪里见过 AI？" in first_slide_text
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in first_slide.shapes)
