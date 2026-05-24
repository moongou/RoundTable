#!/usr/bin/env python3
"""从讲稿 markdown 生成可编辑文本框版 PPTX。"""

from __future__ import annotations

import argparse
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "docs" / "小学四年级人工智能奇妙之旅_讲稿.md"
DEFAULT_OUTPUT = REPO_ROOT / "docs" / "小学四年级人工智能奇妙之旅.pptx"
HEADING_RE = re.compile(r"^## 第(?P<number>\d+)页｜(?P<title>.+)$", re.MULTILINE)
FIELD_RE = re.compile(r"^\*\*(?P<label>[^*：:]+)[：:]\*\*\s*(?P<inline>.*)$")

FONT_NAME = "PingFang SC"
COLOR_BG = RGBColor(245, 248, 255)
COLOR_PANEL = RGBColor(255, 255, 255)
COLOR_PANEL_ALT = RGBColor(237, 244, 255)
COLOR_ACCENT = RGBColor(59, 130, 246)
COLOR_ACCENT_DARK = RGBColor(30, 64, 175)
COLOR_TEXT = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(75, 85, 99)
COLOR_WARM = RGBColor(249, 115, 22)
COLOR_DIVIDER = RGBColor(191, 219, 254)


@dataclass(frozen=True)
class LessonSlide:
    number: int
    title: str
    positioning: str
    script: str
    prompt: str | None
    goals: list[str]


def parse_lesson_markdown(path: Path) -> list[LessonSlide]:
    text = path.read_text(encoding="utf-8")
    matches = list(HEADING_RE.finditer(text))
    slides: list[LessonSlide] = []

    for index, match in enumerate(matches):
        block_start = match.end()
        block_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[block_start:block_end].strip()
        slide = _parse_slide_block(
            int(match.group("number")),
            match.group("title").strip(),
            block,
        )
        slides.append(slide)

    if not slides:
        raise ValueError(f"未在 {path} 中找到任何“## 第N页｜标题”格式的页面定义")

    return slides


def _parse_slide_block(number: int, title: str, block: str) -> LessonSlide:
    fields: dict[str, list[str]] = {}
    lines = block.splitlines()
    index = 0

    while index < len(lines):
        match = FIELD_RE.match(lines[index].strip())
        if not match:
            index += 1
            continue

        label = match.group("label").strip()
        values: list[str] = []
        inline = match.group("inline").strip()
        if inline:
            values.append(inline)
        index += 1

        while index < len(lines) and not FIELD_RE.match(lines[index].strip()):
            values.append(lines[index].rstrip())
            index += 1

        fields[label] = values

    positioning = _join_paragraphs(fields.get("页面定位", []))
    script = _join_paragraphs(fields.get("讲稿", []))
    prompt = _join_paragraphs(fields.get("互动提问", [])) or None
    goals = _extract_bullets(fields.get("本页目标", []))

    return LessonSlide(
        number=number,
        title=title,
        positioning=positioning,
        script=script,
        prompt=prompt,
        goals=goals,
    )


def _join_paragraphs(lines: Iterable[str]) -> str:
    paragraphs: list[str] = []
    current: list[str] = []

    for raw in lines:
        line = raw.strip()
        if not line:
            if current:
                paragraphs.append("".join(current))
                current = []
            continue
        current.append(line)

    if current:
        paragraphs.append("".join(current))

    return "\n\n".join(paragraphs).strip()


def _extract_bullets(lines: Iterable[str]) -> list[str]:
    bullets: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        bullets.append(line[2:].strip() if line.startswith("- ") else line)
    return bullets


def build_presentation(slides: list[LessonSlide], output_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for slide_spec in slides:
        _add_lesson_slide(presentation, slide_spec)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def _add_lesson_slide(presentation: Presentation, slide_spec: LessonSlide) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, COLOR_BG)
    _add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.2, COLOR_ACCENT_DARK)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 2.0, 8.45, 4.8, COLOR_PANEL)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.35, 2.0, 3.38, 2.35, COLOR_PANEL_ALT)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.35, 4.55, 3.38, 2.25, COLOR_PANEL)
    _add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 11.72, 0.38, 1.0, 0.56, COLOR_ACCENT)

    _add_title(slide, slide_spec)
    _add_positioning(slide, slide_spec.positioning)
    _add_script_panel(slide, slide_spec.script)
    _add_goals_panel(slide, slide_spec.goals)
    _add_prompt_panel(slide, slide_spec.prompt)
    _add_slide_number(slide, slide_spec.number)
    _add_footer(slide)


def _add_shape(
    slide,
    shape_type: MSO_SHAPE,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    return shape


def _add_title(slide, slide_spec: LessonSlide) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.42), Inches(10.7), Inches(0.62))
    frame = title_box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = slide_spec.title
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.line_spacing = 1.0
    _style_run(paragraph, Pt(25), COLOR_ACCENT_DARK, bold=True)


def _add_positioning(slide, positioning: str) -> None:
    ribbon = _add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        0.6,
        1.16,
        6.9,
        0.5,
        COLOR_PANEL_ALT,
    )
    ribbon.line.color.rgb = COLOR_DIVIDER

    label_box = slide.shapes.add_textbox(Inches(0.82), Inches(1.26), Inches(6.4), Inches(0.28))
    frame = label_box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = f"页面定位｜{positioning}"
    paragraph.alignment = PP_ALIGN.LEFT
    _style_run(paragraph, Pt(14), COLOR_MUTED, bold=True)


def _add_script_panel(slide, script: str) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.92), Inches(2.18), Inches(1.5), Inches(0.3))
    title_frame = title_box.text_frame
    title_frame.margin_left = 0
    title_frame.margin_right = 0
    title_frame.margin_top = 0
    title_frame.margin_bottom = 0
    paragraph = title_frame.paragraphs[0]
    paragraph.text = "讲解内容"
    _style_run(paragraph, Pt(15), COLOR_ACCENT_DARK, bold=True)

    body_box = slide.shapes.add_textbox(Inches(0.92), Inches(2.58), Inches(7.82), Inches(3.86))
    frame = body_box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)

    paragraphs = [item.strip() for item in script.split("\n\n") if item.strip()]
    font_size = _script_font_size(script)

    first = True
    for block in paragraphs:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        paragraph.text = block
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.25
        paragraph.space_after = Pt(7)
        _style_run(paragraph, font_size, COLOR_TEXT)
        first = False


def _add_goals_panel(slide, goals: list[str]) -> None:
    title_box = slide.shapes.add_textbox(Inches(9.66), Inches(2.18), Inches(1.6), Inches(0.3))
    title_frame = title_box.text_frame
    title_frame.margin_left = 0
    title_frame.margin_right = 0
    title_frame.margin_top = 0
    title_frame.margin_bottom = 0
    paragraph = title_frame.paragraphs[0]
    paragraph.text = "本页目标"
    _style_run(paragraph, Pt(15), COLOR_ACCENT_DARK, bold=True)

    body_box = slide.shapes.add_textbox(Inches(9.66), Inches(2.56), Inches(2.72), Inches(1.5))
    frame = body_box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0

    items = goals or ["可按授课需要补充本页目标。"]
    for index, goal in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {goal}"
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(8)
        _style_run(paragraph, Pt(16), COLOR_TEXT)


def _add_prompt_panel(slide, prompt: str | None) -> None:
    title_box = slide.shapes.add_textbox(Inches(9.66), Inches(4.74), Inches(1.6), Inches(0.3))
    title_frame = title_box.text_frame
    title_frame.margin_left = 0
    title_frame.margin_right = 0
    title_frame.margin_top = 0
    title_frame.margin_bottom = 0
    paragraph = title_frame.paragraphs[0]
    paragraph.text = "互动提问"
    _style_run(paragraph, Pt(15), COLOR_WARM, bold=True)

    body_box = slide.shapes.add_textbox(Inches(9.66), Inches(5.12), Inches(2.72), Inches(1.22))
    frame = body_box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0

    paragraph = frame.paragraphs[0]
    paragraph.text = prompt or "本页没有预设提问，可按课堂氛围灵活发挥。"
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.line_spacing = 1.2
    _style_run(paragraph, Pt(15), COLOR_TEXT)


def _add_slide_number(slide, number: int) -> None:
    box = slide.shapes.add_textbox(Inches(11.98), Inches(0.48), Inches(0.48), Inches(0.22))
    frame = box.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = str(number)
    paragraph.alignment = PP_ALIGN.CENTER
    _style_run(paragraph, Pt(14), RGBColor(255, 255, 255), bold=True)


def _add_footer(slide) -> None:
    footer = slide.shapes.add_textbox(Inches(0.72), Inches(6.98), Inches(8.8), Inches(0.24))
    frame = footer.text_frame
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = "根据讲稿自动生成，全部内容均为可编辑文本框。"
    paragraph.alignment = PP_ALIGN.LEFT
    _style_run(paragraph, Pt(10.5), COLOR_MUTED)


def _style_run(paragraph, size: Pt, color: RGBColor, *, bold: bool = False) -> None:
    if not paragraph.runs:
        paragraph.text = paragraph.text or ""
    run = paragraph.runs[0]
    font = run.font
    font.name = FONT_NAME
    font.size = size
    font.bold = bold
    font.color.rgb = color


def _script_font_size(script: str) -> Pt:
    length = len(script)
    if length > 320:
        return Pt(12.5)
    if length > 260:
        return Pt(13.5)
    if length > 220:
        return Pt(14.5)
    if length > 180:
        return Pt(15.5)
    return Pt(16.5)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="讲稿 markdown 路径")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="输出 pptx 路径")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    slides = parse_lesson_markdown(args.source)
    build_presentation(slides, args.output)
    print(f"Generated editable PPTX: {args.output}")


if __name__ == "__main__":
    main()
